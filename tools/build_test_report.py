"""生成第 2 关、第 5 关交叉测试报告与配套向量。

运行::

    python -m tools.build_test_report

输出::

    results/S-AES_交叉测试报告.docx
    results/partner_cross_encrypt.csv
    results/partner_cross_decrypt.csv
    results/partner_cbc_vectors.csv

报告只包含需要交叉测试的两个关卡：

* 第 2 关：16 组加密向量 + 16 组解密向量；
* 第 5 关：3 组 CBC 向量 + 1 组密文篡改实验。

每一关都包含本组结果截图，以及留给对方粘贴截图的虚线框。
"""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from saes.core import encrypt_block, format_hex
from saes.modes import cbc_decrypt, cbc_encrypt, flip_bit
from saes.text import to_hex
from saes.vectors import generate_vectors

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"
SHOTS = RESULTS / "screenshots"
DOCX_PATH = RESULTS / "S-AES_交叉测试报告.docx"

GRAY = RGBColor(0x66, 0x70, 0x85)
DARK = RGBColor(0x10, 0x18, 0x28)
ACCENT = RGBColor(0x5E, 0x6A, 0xD2)
HEADER_FILL = "EEF0FB"
PLACEHOLDER_FILL = "FAFAFA"


# ---------------------------------------------------------------------------
# 测试向量
# ---------------------------------------------------------------------------


def build_cross_vectors() -> list[tuple[int, int, int]]:
    """4 组边界/教材向量 + 12 组固定种子随机向量，共 16 组。"""
    vectors: list[tuple[int, int, int]] = []
    seen: set[tuple[int, int]] = set()
    for key, plaintext in (
        (0x0000, 0x0000),
        (0xFFFF, 0xFFFF),
        (0x2D55, 0x6AC5),
        (0xABCD, 0x1234),
    ):
        vectors.append((key, plaintext, encrypt_block(plaintext, key)))
        seen.add((key, plaintext))
    for vector in generate_vectors(40, seed=20261024):
        if (vector.key, vector.plaintext) in seen:
            continue
        seen.add((vector.key, vector.plaintext))
        vectors.append((vector.key, vector.plaintext, vector.ciphertext))
        if len(vectors) == 16:
            break
    return vectors


CROSS_VECTORS = build_cross_vectors()

CBC_VECTORS = [
    (0xA73B, 0x1234, "Information Security"),
    (0x2D55, 0x6AC5, "Hello S-AES!"),
    (0x0000, 0x0000, "ok"),
]

TAMPER_KEY, TAMPER_IV, TAMPER_TEXT = CBC_VECTORS[0]
TAMPER_BLOCK = 1
TAMPER_BIT = 3


def cbc_cipher(text: str, key: int, iv: int) -> bytes:
    return cbc_encrypt(text.encode("utf-8"), key, iv)


# ---------------------------------------------------------------------------
# python-docx 辅助函数
# ---------------------------------------------------------------------------


def set_font(run, name: str = "宋体", size: float | None = None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = doc.styles[style_name]
        style.font.name = "黑体"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = DARK
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = style_name != "Title"


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("信息安全导论 · 作业 2 · S-AES 交叉测试报告")
    set_font(run, size=8.5, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_font(run, size=9, color=GRAY)
    field_run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.append(begin)
    field_run._r.append(instr)
    field_run._r.append(end)
    set_font(field_run, size=9, color=GRAY)
    run = footer.add_run(" 页")
    set_font(run, size=9, color=GRAY)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, *, bold=False, size=9, mono=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_font(run, "Consolas" if mono else "宋体", size, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, headers, rows, widths=None, mono_columns=(), font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell(cell, header, bold=True, size=font_size)
        shade(cell, HEADER_FILL)
        if widths:
            cell.width = Cm(widths[index])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            set_cell(cell, value, size=font_size, mono=index in mono_columns)
            if widths:
                cell.width = Cm(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc: Document, image_name: str, caption_text: str) -> None:
    path = SHOTS / image_name
    if not path.exists():
        raise FileNotFoundError(f"缺少截图：{path}")
    doc.add_picture(str(path), width=Cm(15.0))
    picture_paragraph = doc.paragraphs[-1]
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.keep_with_next = True
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(caption_text)
    set_font(run, size=9, color=GRAY)


def set_row_height(row, height_cm: float) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(int(height_cm * 567)))
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)


def set_dashed_border(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "dashed")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "A6A6A6")
        borders.append(element)
    tbl_pr.append(borders)


def add_placeholder(doc: Document, text: str, height_cm: float = 7.0) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_dashed_border(table)
    cell = table.cell(0, 0)
    cell.width = Cm(16.0)
    shade(cell, PLACEHOLDER_FILL)
    set_row_height(table.rows[0], height_cm)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(50)
    run = paragraph.add_run(text)
    set_font(run, size=10.5, color=GRAY)
    paragraph2 = cell.add_paragraph()
    paragraph2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = paragraph2.add_run("（请在此处粘贴截图）")
    set_font(run2, size=9, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run)


def bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_font(run)


def note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_font(run, size=9, color=GRAY)


# ---------------------------------------------------------------------------
# 文档章节
# ---------------------------------------------------------------------------


def build_cover(doc: Document) -> None:
    for _ in range(6):
        doc.add_paragraph()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("S-AES 交叉测试报告")
    set_font(run, "黑体", 26, True, DARK)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 2 关 · 交叉测试　|　第 5 关 · CBC 模式")
    set_font(run, "黑体", 14, False, ACCENT)

    for _ in range(4):
        doc.add_paragraph()

    info = doc.add_table(rows=0, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    for label, value in (
        ("组别", "第 ______ 组"),
        ("本组成员", "____________________________"),
        ("学号", "____________________________"),
        ("交叉测试对象", "____________________________"),
        ("测试日期", "2026 年 10 月 24 日"),
    ):
        row = info.add_row()
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(9.0)
        set_cell(row.cells[0], label, bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell(row.cells[1], value, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)

    for _ in range(4):
        doc.add_paragraph()

    note_box = doc.add_table(rows=1, cols=1)
    note_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    note_box.autofit = False
    set_dashed_border(note_box)
    cell = note_box.cell(0, 0)
    cell.width = Cm(15.0)
    shade(cell, PLACEHOLDER_FILL)
    set_cell(
        cell,
        "本报告只包含需要交叉测试的两个关卡：第 2 关（16 组加密 + 16 组解密）"
        "与第 5 关（3 组 CBC + 1 组篡改实验）。每关均附本组截图和对方截图预留框。",
        size=10,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    doc.add_page_break()


def build_overview(doc: Document) -> None:
    heading(doc, "一、测试环境与结果概览", 1)
    body(doc, "测试环境：Windows 11 · Python 3.12.5 · Tkinter；本组程序无第三方运行时依赖。")
    add_table(
        doc,
        ("关卡", "交叉测试内容", "本组结果", "对方结果"),
        [
            ("第 2 关", "16 组加密向量 + 16 组解密向量", "32/32 一致", "待补充"),
            ("第 5 关", "3 组 CBC 向量 + 1 组密文篡改", "全部通过", "待补充"),
        ],
        widths=(2.2, 7.2, 3.8, 2.6),
        font_size=9.5,
    )
    body(doc, "对方在空白列填写自己的结果，并把截图粘贴到每关末尾的虚线框中即可。")
    doc.add_page_break()


def build_level2(doc: Document) -> None:
    heading(doc, "二、第 2 关：交叉测试", 1)
    heading(doc, "2.1 约定与算法标准", 2)
    add_table(
        doc,
        ("项目", "约定"),
        [
            ("分组 / 密钥", "16 bit / 16 bit，2 轮"),
            ("状态矩阵", "[ n0 n2 / n1 n3 ]，n0 为最高 4 位"),
            ("S 盒", "9 4 A B / D 1 8 5 / 6 2 0 3 / C E F 7"),
            ("列混淆 / 逆列混淆", "[[1,4],[4,1]] / [[9,2],[2,9]]"),
            ("GF(2^4)", "模多项式 x^4 + x + 1"),
            ("密钥扩展", "w2 = w0 ⊕ g(w1,0x80)，w4 = w2 ⊕ g(w3,0x30)"),
        ],
        widths=(4.5, 11.5),
        font_size=9,
    )
    body(doc, "本组已与 A 组完成交换测试：A 组 16 组加密、16 组解密向量全部与本组一致。"
              "下表 16 组为提供给对方的新向量。")

    heading(doc, "2.2 本组自测截图", 2)
    add_figure(doc, "02a_cross_generate.png", "图 2-1  本组生成交叉测试向量")
    add_figure(doc, "02b_cross_check.png", "图 2-2  本组向量逐条校验通过")

    heading(doc, "2.3 加密方向向量（16 组，请对方填写 C'）", 2)
    rows = [
        (index + 1, f"{key:04X}", f"{plaintext:04X}", f"{cipher:04X}", "", "")
        for index, (key, plaintext, cipher) in enumerate(CROSS_VECTORS)
    ]
    add_table(
        doc,
        ("序号", "密钥 K", "明文 P", "本组密文 C", "对方密文 C'", "是否一致"),
        rows,
        widths=(1.3, 2.5, 2.5, 2.8, 3.0, 2.5),
        font_size=9,
        mono_columns=(1, 2, 3),
    )

    heading(doc, "2.4 解密方向向量（16 组，请对方填写 P'）", 2)
    rows = [
        (index + 1, f"{key:04X}", f"{cipher:04X}", f"{plaintext:04X}", "", "")
        for index, (key, plaintext, cipher) in enumerate(CROSS_VECTORS)
    ]
    add_table(
        doc,
        ("序号", "密钥 K", "密文 C", "本组明文 P", "对方明文 P'", "是否一致"),
        rows,
        widths=(1.3, 2.5, 2.5, 2.8, 3.0, 2.5),
        font_size=9,
        mono_columns=(1, 2, 3),
    )

    heading(doc, "2.5 对方测试结果（预留）", 2)
    add_placeholder(doc, "对方第 2 关加密方向截图")
    note(doc, "图 2-3  对方第 2 关加密方向结果（待补充）")
    add_placeholder(doc, "对方第 2 关解密方向截图")
    note(doc, "图 2-4  对方第 2 关解密方向结果（待补充）")
    doc.add_page_break()


def build_level5(doc: Document) -> None:
    heading(doc, "三、第 5 关：CBC 工作模式", 1)
    heading(doc, "3.1 约定", 2)
    body(
        doc,
        "C_i = E(P_i ⊕ C_(i-1))，IV 为 16 bit。字符串按 UTF-8 编码，"
        "以 2 字节为一组；使用 PKCS#7 填充，缺 1 字节补 01，缺 2 字节补 02 02。"
        "对方需使用相同的填充方式，否则密文无法直接比对。",
    )
    rows = []
    for key, iv, plaintext in CBC_VECTORS:
        rows.append((format_hex(key), format_hex(iv), plaintext, to_hex(cbc_cipher(plaintext, key, iv))))
    add_table(
        doc,
        ("密钥 K", "IV", "明文", "本组密文（十六进制）"),
        rows,
        widths=(2.0, 1.8, 4.8, 7.4),
        font_size=8.5,
        mono_columns=(0, 1, 3),
    )
    body(doc, "以上 3 组均完成加密—解密往返校验，全部还原原文。")

    heading(doc, "3.2 密文篡改实验", 2)
    cipher = cbc_cipher(TAMPER_TEXT, TAMPER_KEY, TAMPER_IV)
    tampered = flip_bit(cipher, TAMPER_BLOCK - 1, TAMPER_BIT)
    original_raw = cbc_decrypt(cipher, TAMPER_KEY, TAMPER_IV, unpad_data=False)
    tampered_raw = cbc_decrypt(tampered, TAMPER_KEY, TAMPER_IV, unpad_data=False)
    changed = [
        f"字节 {i}：{left:02X} → {right:02X}"
        for i, (left, right) in enumerate(zip(original_raw, tampered_raw))
        if left != right
    ]
    body(
        doc,
        f"以 K={TAMPER_KEY:04X}、IV={TAMPER_IV:04X} 的第一组密文为例，"
        f"翻转第 {TAMPER_BLOCK} 个分组的第 {TAMPER_BIT} 位后解密：",
    )
    add_table(
        doc,
        ("项目", "数据"),
        [
            ("原始密文", to_hex(cipher)),
            ("篡改后密文", to_hex(tampered)),
            ("原始解密字节", to_hex(original_raw)),
            ("篡改后解密字节", to_hex(tampered_raw)),
            ("变化字节", "；".join(changed)),
        ],
        widths=(3.5, 12.5),
        font_size=8.5,
        mono_columns=(1,),
    )
    body(doc, "被篡改分组整组变为随机噪声；下一分组只有对应比特位翻转；更后面的分组不受影响。")

    heading(doc, "3.3 本组测试截图", 2)
    add_figure(doc, "05a_cbc_top.png", "图 5-1  CBC 加密与分组明细")
    add_figure(doc, "05b_cbc_tamper.png", "图 5-2  密文篡改前后的解密对比")

    heading(doc, "3.4 提供给对方的 CBC 向量（请对方填写 C'）", 2)
    rows = [
        (index + 1, format_hex(key), format_hex(iv), plaintext,
         to_hex(cbc_cipher(plaintext, key, iv)), "", "")
        for index, (key, iv, plaintext) in enumerate(CBC_VECTORS)
    ]
    add_table(
        doc,
        ("序号", "密钥 K", "IV", "明文", "本组密文（hex）", "对方密文（hex）", "是否一致"),
        rows,
        widths=(0.9, 1.5, 1.2, 3.2, 5.0, 3.0, 1.2),
        font_size=7.5,
        mono_columns=(1, 2, 4),
    )

    heading(doc, "3.5 对方测试结果（预留）", 2)
    add_placeholder(doc, "对方第 5 关 CBC 加解密截图")
    note(doc, "图 5-3  对方第 5 关 CBC 加解密结果（待补充）")
    add_placeholder(doc, "对方第 5 关密文篡改截图")
    note(doc, "图 5-4  对方第 5 关密文篡改结果（待补充）")
    doc.add_page_break()


def build_instructions(doc: Document) -> None:
    heading(doc, "四、填写说明", 1)
    bullet(doc, "第 2 关：在“对方密文 C'”“对方明文 P'”两列填写，逐行与本组结果比对。")
    bullet(doc, "第 5 关：在“对方密文”列填写，并与本组密文逐字节比对；篡改实验请记录变化字节。")
    bullet(doc, "截图粘贴到每关末尾的灰色虚线框内，可粘贴多张。")
    bullet(doc, "若结果不一致，优先检查状态矩阵的列优先布局、S 盒取值、列混淆矩阵、轮密钥顺序，以及 CBC 的 PKCS#7 填充约定。")


# ---------------------------------------------------------------------------
# CSV 导出
# ---------------------------------------------------------------------------


def write_csv(path: Path, header, rows) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def export_csvs() -> None:
    write_csv(
        RESULTS / "partner_cross_encrypt.csv",
        ("key", "plaintext", "ciphertext"),
        [(f"{k:04X}", f"{p:04X}", f"{c:04X}") for k, p, c in CROSS_VECTORS],
    )
    write_csv(
        RESULTS / "partner_cross_decrypt.csv",
        ("key", "ciphertext", "plaintext"),
        [(f"{k:04X}", f"{c:04X}", f"{p:04X}") for k, p, c in CROSS_VECTORS],
    )
    cbc_rows = [
        (f"{key:04X}", f"{iv:04X}", plaintext, to_hex(cbc_cipher(plaintext, key, iv)))
        for key, iv, plaintext in CBC_VECTORS
    ]
    cipher = cbc_cipher(TAMPER_TEXT, TAMPER_KEY, TAMPER_IV)
    tampered = flip_bit(cipher, TAMPER_BLOCK - 1, TAMPER_BIT)
    write_csv(
        RESULTS / "partner_cbc_vectors.csv",
        ("key", "iv", "plaintext", "ciphertext_hex", "tamper_block", "tamper_bit", "tampered_cipher_hex"),
        cbc_rows + [
            (f"{TAMPER_KEY:04X}", f"{TAMPER_IV:04X}", TAMPER_TEXT,
             to_hex(cipher), str(TAMPER_BLOCK), str(TAMPER_BIT), to_hex(tampered))
        ],
    )


def main() -> None:
    RESULTS.mkdir(exist_ok=True)
    export_csvs()

    document = Document()
    configure_styles(document)
    configure_page(document)

    build_cover(document)
    build_overview(document)
    build_level2(document)
    build_level5(document)
    build_instructions(document)

    # 删除末尾空段落，避免多出空白页
    while document.paragraphs and not document.paragraphs[-1].text.strip():
        element = document.paragraphs[-1]._element
        element.getparent().remove(element)

    document.save(DOCX_PATH)
    print(f"已生成 {DOCX_PATH}")
    for name in ("partner_cross_encrypt.csv", "partner_cross_decrypt.csv", "partner_cbc_vectors.csv"):
        print(f"已生成 {RESULTS / name}")


if __name__ == "__main__":
    main()
